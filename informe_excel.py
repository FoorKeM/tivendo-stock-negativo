import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from tkinter import Tk, filedialog

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from analizar_inventario import AdjustmentItem
from app_paths import ADJUSTMENT_DIR


HEADERS = (
    "Fecha ajuste",
    "Bodega",
    "CodArticulo",
    "Descripcion Articulo",
    "Codigo barra interno",
    "Stock negativo",
    "Cantidad ajustada",
)


def _number(value: str):
    decimal = Decimal(str(value).strip().replace(",", "."))
    if decimal == decimal.to_integral():
        return int(decimal)
    return float(decimal)


def _default_name() -> str:
    return f"Informe_Stock_Negativo_{datetime.now():%Y%m%d_%H%M%S}.xlsx"


def elegir_destino() -> Path | None:
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    try:
        selected = filedialog.asksaveasfilename(
            parent=root,
            title="Guardar informe de stock negativo",
            defaultextension=".xlsx",
            initialfile=_default_name(),
            filetypes=[("Libro de Excel", "*.xlsx")],
        )
    finally:
        root.destroy()
    return Path(selected) if selected else None


def crear_informe(
    result_path: Path,
    items: list[AdjustmentItem],
    destination: Path,
) -> Path:
    payload = json.loads(result_path.read_text(encoding="utf-8"))
    descriptions = {item.codigo: item.descripcion for item in items}
    rows = []

    for article in payload.get("articulos", []):
        if article.get("estado") != "cargado":
            continue
        stock = Decimal(str(article["stock_actual"]).strip().replace(",", "."))
        quantity = Decimal(str(article["cantidad"]).strip().replace(",", "."))
        if stock >= 0 or quantity <= 0:
            continue
        code = str(article["codigo"]).strip()
        if not code.upper().startswith("A"):
            continue
        rows.append(
            [
                payload["fecha"],
                payload["bodega"],
                code,
                descriptions.get(code, ""),
                str(article["codigo_barra"]).strip(),
                _number(str(stock)),
                _number(str(quantity)),
            ]
        )

    if not rows:
        raise RuntimeError("No hay articulos negativos ajustados para el informe.")

    destination = destination.with_suffix(".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stock negativo"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)

    header = sheet[1]
    for cell in header:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for cell in sheet["E"][1:]:
        cell.number_format = "@"
    for row in sheet.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = "0.########"

    widths = {
        "A": 22,
        "B": 20,
        "C": 16,
        "D": 55,
        "E": 24,
        "F": 16,
        "G": 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A2"
    table = Table(displayName="AjustesStockNegativo", ref=f"A1:G{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    workbook.save(destination)
    workbook.close()

    check = load_workbook(destination, read_only=True, data_only=True)
    try:
        if check.active.max_row != len(rows) + 1:
            raise RuntimeError("El informe Excel no contiene todas las filas esperadas.")
    finally:
        check.close()
    return destination


def guardar_informe_interactivo(
    result_path: Path,
    items: list[AdjustmentItem],
) -> Path:
    destination = elegir_destino()
    if destination is None:
        destination = ADJUSTMENT_DIR / _default_name()
        print(f"No se eligio una carpeta. Se guardara en: {destination}")
    return crear_informe(result_path, items, destination)
