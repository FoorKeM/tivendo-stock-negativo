import json
import tempfile
import zipfile
from dataclasses import asdict, dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from app_paths import ADJUSTMENT_DIR


INVENTORY_HEADERS = ("CodArticulo", "Descripción Artículo", "Cantidad")
ARTICLE_HEADERS = (
    "Código",
    "Nombre",
    "Código barra interno",
    "Disponible para venta",
    "Activo",
)


@dataclass(frozen=True)
class AdjustmentItem:
    codigo: str
    descripcion: str
    stock_informe: str
    codigo_barra: str
    cantidad_propuesta: str


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _decimal(value) -> Decimal:
    if isinstance(value, Decimal):
        return value
    text = _text(value).replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Cantidad invalida: {value!r}") from exc


def _decimal_text(value: Decimal) -> str:
    return format(value, "f")


def _read_rows(path: Path, required_headers: tuple[str, ...]):
    repaired_path = _repair_null_numbers(path) if _has_null_numbers(path) else None
    workbook = load_workbook(
        repaired_path or path,
        read_only=True,
        data_only=True,
    )
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        for row in iterator:
            normalized = tuple(_text(value) for value in row)
            if all(header in normalized for header in required_headers):
                indexes = {header: normalized.index(header) for header in required_headers}
                return indexes, list(iterator)
    finally:
        workbook.close()
        if repaired_path:
            repaired_path.unlink(missing_ok=True)
    raise ValueError(
        f"No se encontraron las columnas requeridas en {path.name}: "
        + ", ".join(required_headers)
    )


def _has_null_numbers(path: Path) -> bool:
    with zipfile.ZipFile(path, "r") as source:
        for name in source.namelist():
            if name.startswith("xl/worksheets/"):
                if b"<v>null</v>" in source.read(name):
                    return True
    return False


def _repair_null_numbers(path: Path) -> Path:
    handle = tempfile.NamedTemporaryFile(
        prefix="tivendo_reparado_",
        suffix=".xlsx",
        delete=False,
    )
    repaired = Path(handle.name)
    handle.close()
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        repaired,
        "w",
        zipfile.ZIP_DEFLATED,
    ) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/"):
                data = data.replace(b"<v>null</v>", b"<v></v>")
            target.writestr(info, data)
    return repaired


def leer_negativos(path: Path) -> list[dict]:
    indexes, rows = _read_rows(path, INVENTORY_HEADERS)
    negatives = []
    for row in rows:
        code = _text(row[indexes["CodArticulo"]])
        if not code.upper().startswith("A"):
            continue
        quantity = _decimal(row[indexes["Cantidad"]])
        if quantity >= 0:
            continue
        negatives.append(
            {
                "codigo": code,
                "descripcion": _text(row[indexes["Descripción Artículo"]]),
                "stock": quantity,
            }
        )
    negatives.sort(key=lambda item: (item["stock"], item["codigo"]))
    return negatives


def leer_articulos(path: Path) -> dict[str, list[dict]]:
    indexes, rows = _read_rows(path, ARTICLE_HEADERS)
    articles: dict[str, list[dict]] = {}
    for row in rows:
        code = _text(row[indexes["Código"]])
        if not code:
            continue
        articles.setdefault(code, []).append(
            {
                "nombre": _text(row[indexes["Nombre"]]),
                "codigo_barra": _text(row[indexes["Código barra interno"]]),
                "venta": _text(row[indexes["Disponible para venta"]]),
                "activo": _text(row[indexes["Activo"]]),
            }
        )
    return articles


def cruzar_negativos(inventory_path: Path, articles_path: Path) -> list[AdjustmentItem]:
    negatives = leer_negativos(inventory_path)
    articles = leer_articulos(articles_path)
    errors = []
    result = []

    for negative in negatives:
        matches = articles.get(negative["codigo"], [])
        if len(matches) != 1:
            errors.append(
                f"{negative['codigo']}: se esperaban 1 coincidencia y hay {len(matches)}"
            )
            continue
        article = matches[0]
        if not article["codigo_barra"]:
            errors.append(f"{negative['codigo']}: sin Codigo barra interno")
            continue
        if article["activo"].casefold() != "si":
            errors.append(f"{negative['codigo']}: articulo inactivo")
            continue
        if article["venta"].casefold() != "si":
            errors.append(f"{negative['codigo']}: no disponible para venta")
            continue

        quantity = abs(negative["stock"])
        result.append(
            AdjustmentItem(
                codigo=negative["codigo"],
                descripcion=negative["descripcion"],
                stock_informe=_decimal_text(negative["stock"]),
                codigo_barra=article["codigo_barra"],
                cantidad_propuesta=_decimal_text(quantity),
            )
        )

    if errors:
        raise ValueError("No se puede preparar el ajuste:\n- " + "\n- ".join(errors))
    return result


def guardar_propuesta(
    inventory_path: Path,
    articles_path: Path,
    items: list[AdjustmentItem],
) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    destination = ADJUSTMENT_DIR / f"propuesta_ajuste_{stamp}.json"
    payload = {
        "creada": datetime.now().isoformat(timespec="seconds"),
        "bodega": "LIBERTADOR 1476",
        "inventario": str(inventory_path),
        "listado_articulos": str(articles_path),
        "cantidad_articulos": len(items),
        "articulos": [asdict(item) for item in items],
    }
    destination.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return destination
