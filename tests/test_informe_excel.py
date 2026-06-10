import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from analizar_inventario import AdjustmentItem
from informe_excel import HEADERS, crear_informe


class NegativeStockReportTests(unittest.TestCase):
    def test_informe_solo_contiene_negativos_cargados(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            result = folder / "resultado.json"
            destination = folder / "informe.xlsx"
            result.write_text(
                json.dumps(
                    {
                        "fecha": "2026-06-09T16:34:40",
                        "bodega": "LIBERTADOR 1476",
                        "articulos": [
                            {
                                "codigo": "A000001",
                                "codigo_barra": "0012345",
                                "stock_actual": "-15.5",
                                "cantidad": "15.5",
                                "estado": "cargado",
                            },
                            {
                                "codigo": "A000002",
                                "codigo_barra": "999",
                                "stock_actual": "2",
                                "estado": "omitido_no_negativo",
                            },
                            {
                                "codigo": "P000003",
                                "codigo_barra": "888",
                                "stock_actual": "-3",
                                "cantidad": "3",
                                "estado": "cargado",
                            },
                        ],
                    }
                ),
                encoding="utf-8",
            )
            items = [
                AdjustmentItem(
                    codigo="A000001",
                    descripcion="Articulo de prueba",
                    stock_informe="-15.5",
                    codigo_barra="0012345",
                    cantidad_propuesta="15.5",
                )
            ]

            crear_informe(result, items, destination)

            workbook = load_workbook(destination, read_only=True, data_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()

        self.assertEqual(rows[0], HEADERS)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][2], "A000001")
        self.assertEqual(rows[1][3], "Articulo de prueba")
        self.assertEqual(rows[1][4], "0012345")
        self.assertEqual(rows[1][5], -15.5)
        self.assertEqual(rows[1][6], 15.5)


if __name__ == "__main__":
    unittest.main()
